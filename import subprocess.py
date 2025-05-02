import subprocess
import time
import openpyxl
import re
import logging
import argparse
import os
import pathlib
import platform
import ipaddress
from datetime import datetime

# Logging setup
logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)
formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(filename)s:%(lineno)d - %(message)s')

# Ping 次數以及重試次數
ping_count = 4
retry_attempts = 3

def open_or_create_workbook(excel_file):
    """開啟現有或建立新的 Excel 工作簿，並處理各種可能的例外狀況"""
    try:
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
        # 檢查工作表是否存在且第一列非空
        if sheet.title != "Sheet1" or not sheet["A1"].value:
            if "Sheet1" in workbook.sheetnames:
                sheet = workbook["Sheet1"]
                if not sheet["A1"].value:  # 只有在Sheet1存在且為空時才新增標題列
                    sheet.append(["時間", "最小值(ms)", "最大值(ms)", "平均值(ms)", "遺失率(%)"])
            else:
                sheet = workbook.create_sheet("Sheet1", 0)
                sheet.append(["時間", "最小值(ms)", "最大值(ms)", "平均值(ms)", "遺失率(%)"])
        return workbook, sheet
    except FileNotFoundError:
        logging.info(f"Excel檔案 '{excel_file}' 不存在，將建立新的檔案。")  # 更友善的訊息
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(["時間", "最小值(ms)", "最大值(ms)", "平均值(ms)", "遺失率(%)"])
        return workbook, sheet
    except PermissionError:
        logging.error(f"沒有權限存取 Excel 檔案 '{excel_file}'。")
        return None, None #明確回傳None, None 表示失敗
    except Exception as e:
        logging.exception(f"開啟或建立 Excel 檔案發生錯誤: {e}") # 使用logging.exception 記錄詳細的錯誤訊息
        return None, None


def ping_and_write_to_excel(target_ip, excel_file, ping_packet_count, retry_attempts):
    """執行 ping 指令並將結果寫入 Excel 試算表"""
    workbook, sheet = open_or_create_workbook(excel_file)
    if not workbook: #檢查open_or_create_workbook是否成功
        return False

    try:
        ping_results = run_ping_command(target_ip, ping_packet_count, retry_attempts)
        if ping_results:
            current_time = time.strftime("%Y-%m-%d %H:%M:%S")
            sheet.append([current_time, ping_results['min'], ping_results['max'], ping_results['avg'], ping_results['loss']])
            workbook.save(excel_file)
            logging.info(f"Ping 結果已成功寫入 Excel: {current_time}, 最小:{ping_results['min']}, 最大:{ping_results['max']}, 平均:{ping_results['avg']}, 遺失率:{ping_results['loss']}")
            return True  # 回傳 True 表示成功
        else:
            logging.error(f"Ping 指令執行失敗超過 {retry_attempts} 次，無法取得結果。")
            return False  # 回傳 False 表示失敗
    except Exception as e:
        logging.exception(f"寫入 Excel 發生錯誤: {e}")
        return False  # 回傳 False 表示失敗
    finally:
        if workbook:
            workbook.close()

def run_ping_command(target_ip, ping_count, retries):
    system = platform.system()
    for attempt in range(retries):
        try:
            if system == "Windows":
                command = ['ping', '-n', str(ping_count), target_ip]
            elif system in ["Linux", "Darwin"]:
                command = ['ping', '-c', str(ping_count), target_ip]
            else:
                logging.error(f"Unsupported OS: {system}")
                return None

            process = subprocess.run(command, capture_output=True, text=True, timeout=5, check=True)
            output = process.stdout

            results = parse_ping_output(output, system)
            if results:
                return results
            else:
                logging.warning(f"Ping parsing failed for {target_ip} (Attempt {attempt+1}).")
                return None

        except subprocess.CalledProcessError as e:
            # Added more detailed error logging
            logging.error(f"Ping command failed (attempt {attempt + 1}/{retries}): return code {e.returncode}, output: {e.output}")
            if attempt == retries - 1:
                return None
            time.sleep(2 ** attempt)
        except subprocess.TimeoutExpired:
            logging.error(f"Ping command timed out (attempt {attempt + 1}/{retries})")
            if attempt == retries - 1:
                return None
            time.sleep(2 ** attempt)
        except Exception as e:
            logging.exception(f"Unexpected error during ping: {e}")
            return None
    return None

def parse_ping_output(output, system):
    results = {'min': 'N/A', 'max': 'N/A', 'avg': 'N/A', 'loss': 0}
    logging.debug(f"Ping Output (decoded):\n{output}")

    try:
        times_ms = re.findall(r"(\d+)ms", output) # Corrected regex for string
        if times_ms:
            nums = [int(x) for x in times_ms]
            results['min'] = min(nums)
            results['max'] = max(nums)
            results['avg'] = sum(nums) // len(nums)

        # Packet loss extraction (adjust regex if necessary for your system)
        if system == "Windows":
            loss_match = re.search(r"(\d+)% 遺失", output, re.IGNORECASE) # Corrected regex for string
            if loss_match:
                results['loss'] = int(loss_match.group(1))
        elif system in ["Linux", "Darwin"]:
            loss_match = re.search(r"(\d+)% packet loss", output, re.IGNORECASE) # Corrected regex for string
            if loss_match:
                results['loss'] = int(loss_match.group(1))

    except (AttributeError, IndexError, ValueError) as e:
        logging.warning(f"Error parsing ping output: {e}. Raw Output: {output}")
        return None

    return results

def main():
    parser = argparse.ArgumentParser(description='Ping a target IP and write results to an Excel file.')
    parser.add_argument('--target-ip', help='目標 IP 位址 ', default='10.14.29.201')
    parser.add_argument('--excel-file', help='Excel 檔案路徑 ', default='ping_results.xlsx')
    parser.add_argument('-c', '--count', type=int, default=4, help='Ping 封包數量')
    parser.add_argument('-r', '--retries', type=int, default=3, help='重試次數')
    args = parser.parse_args()

    try:
        ipaddress.ip_address(args.target_ip)
    except ValueError:
        logging.error("無效的 IP 位址")
        return

    today = datetime.now().strftime("%Y%m%d")
    excel_filename = f"{today}_{args.excel_file}"
    log_filename = f"{today}_ping_monitor.log"

    output_dir = pathlib.Path("./ping_results").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    excel_file_path = (output_dir / excel_filename).resolve()
    log_file_path = (output_dir / log_filename).resolve()

    # ALL logging configuration is done here:
    logging.basicConfig(level=logging.DEBUG,
                        format='%(asctime)s - %(levelname)s - %(filename)s:%(lineno)d - %(message)s',
                        filename=str(log_file_path),
                        filemode='w')

    workbook, sheet = None, None
    try:
        print(f"將使用 Excel 檔案: {excel_file_path}")
        print(f"將使用日誌檔案: {log_file_path}")
        workbook, sheet = open_or_create_workbook(str(excel_file_path))
        if not workbook:
            raise Exception("無法開啟或建立 Excel 工作簿")

        while True:
            result = ping_and_write_to_excel(args.target_ip, str(excel_file_path), args.count, args.retries)
            if result:
                print("Ping 結果已成功寫入 Excel。")
            else:
                logger.error("Ping 或寫入 Excel 失敗。請檢查日誌檔案和程式碼。")
                print("Ping 或寫入 Excel 失敗。請檢查日誌檔案和程式碼。")
            time.sleep(60)

    except Exception as e:
        logger.exception(f"主要程式發生錯誤: {e}")
        print(f"主要程式發生錯誤: {e}")
    finally:
        if workbook:
            workbook.close()
            logger.info("Excel 工作簿已關閉")

if __name__ == "__main__":
    main()